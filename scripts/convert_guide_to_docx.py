"""Convert docs/DETECTION_AND_PARAMETERS_GUIDE.md to a beautifully formatted .docx document."""

import re
from pathlib import Path
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_pr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'))

def set_cell_margins(cell, top=120, bottom=120, left=180, right=180):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tc_pr.append(tc_mar)

def set_cell_borders(cell, top="D1D5DB", bottom="D1D5DB", left="D1D5DB", right="D1D5DB"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{top}"/>'
        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="{bottom}"/>'
        f'<w:left w:val="single" w:sz="4" w:space="0" w:color="{left}"/>'
        f'<w:right w:val="single" w:sz="4" w:space="0" w:color="{right}"/>'
        f'</w:tcBorders>'
    )
    tc_pr.append(borders)

def format_inline_text(paragraph, text, default_font="Calibri", default_size=Pt(10.5), default_color=RGBColor(40, 44, 52)):
    # Regular expression for markdown tokens: bold, italic, code, math
    # Tokens: **bold**, `code`, *italic*, $$math$$, $math$
    pattern = re.compile(r'(\*\*.*?\*\*|`.*?`|\*.*?\*|\$\$.*?\$\$|\$.*?\$|__.*?__)')
    parts = pattern.split(text)

    for part in parts:
        if not part:
            continue
        run = paragraph.add_run()
        run.font.name = default_font
        run.font.size = default_size
        run.font.color.rgb = default_color

        if part.startswith('**') and part.endswith('**') and len(part) >= 4:
            run.text = part[2:-2]
            run.bold = True
        elif part.startswith('__') and part.endswith('__') and len(part) >= 4:
            run.text = part[2:-2]
            run.bold = True
        elif part.startswith('`') and part.endswith('`') and len(part) >= 2:
            run.text = part[1:-1]
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(199, 37, 78) # Dark magenta/red for inline code
        elif part.startswith('*') and part.endswith('*') and len(part) >= 2:
            run.text = part[1:-1]
            run.italic = True
        elif (part.startswith('$$') and part.endswith('$$')) or (part.startswith('$') and part.endswith('$')):
            # Clean math symbols
            cleaned = part.strip('$').strip()
            run.text = cleaned
            run.italic = True
            run.font.name = "Cambria Math"
        else:
            run.text = part

def convert_markdown_to_docx(md_path: Path, docx_path: Path):
    doc = docx.Document()

    # Set page margins (1 inch)
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Base styles
    NAVY = RGBColor(16, 44, 87)       # #102C57
    ACCENT_BLUE = RGBColor(30, 86, 160) # #1E56A0
    SLATE = RGBColor(53, 64, 86)      # #354056
    BODY_COLOR = RGBColor(45, 52, 54)  # #2D3436

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_code_block = False
    code_lines = []
    code_lang = ""

    in_table = False
    table_rows = []

    i = 0
    while i < len(lines):
        raw_line = lines[i]
        line = raw_line.rstrip('\r\n')
        stripped = line.strip()

        # 1. Code Block Fence Check
        if stripped.startswith('```'):
            if in_code_block:
                # End code block -> render
                code_text = "\n".join(code_lines)
                table = doc.add_table(rows=1, cols=1)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                cell = table.cell(0, 0)
                set_cell_background(cell, "F4F6F9")
                set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
                set_cell_borders(cell, top="CBD5E1", bottom="CBD5E1", left="CBD5E1", right="CBD5E1")

                p = cell.paragraphs[0]
                p.paragraph_format.line_spacing = 1.05
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                run = p.add_run(code_text)
                run.font.name = "Consolas"
                run.font.size = Pt(8.5)
                run.font.color.rgb = RGBColor(30, 41, 59)

                # Spacing after code block
                sp = doc.add_paragraph()
                sp.paragraph_format.space_before = Pt(0)
                sp.paragraph_format.space_after = Pt(4)

                in_code_block = False
                code_lines = []
                code_lang = ""
            else:
                in_code_block = True
                code_lang = stripped[3:].strip()
                code_lines = []
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # 2. Table Row Check
        if stripped.startswith('|') and stripped.endswith('|'):
            table_rows.append(stripped)
            in_table = True
            i += 1
            continue
        elif in_table:
            # End of table encountered -> Render table
            parsed_rows = []
            for tr in table_rows:
                # Check if it's separator row |---|---|
                cells = [c.strip() for c in tr.split('|')[1:-1]]
                if all(re.match(r'^:?-+:?$', c) for c in cells):
                    continue
                parsed_rows.append(cells)

            if parsed_rows:
                num_cols = max(len(r) for r in parsed_rows)
                tbl = doc.add_table(rows=len(parsed_rows), cols=num_cols)
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER

                for r_idx, row_data in enumerate(parsed_rows):
                    is_header = (r_idx == 0)
                    for c_idx in range(num_cols):
                        c_text = row_data[c_idx] if c_idx < len(row_data) else ""
                        cell = tbl.cell(r_idx, c_idx)
                        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

                        if is_header:
                            set_cell_background(cell, "102C57")
                            set_cell_margins(cell, top=140, bottom=140, left=160, right=160)
                            set_cell_borders(cell, top="0B1E3B", bottom="0B1E3B", left="0B1E3B", right="0B1E3B")
                            p = cell.paragraphs[0]
                            p.paragraph_format.space_before = Pt(2)
                            p.paragraph_format.space_after = Pt(2)
                            format_inline_text(p, c_text, default_size=Pt(9.5), default_color=RGBColor(255, 255, 255))
                            for r in p.runs:
                                r.bold = True
                        else:
                            bg_color = "F9FAFB" if r_idx % 2 == 1 else "FFFFFF"
                            set_cell_background(cell, bg_color)
                            set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
                            set_cell_borders(cell, top="E2E8F0", bottom="E2E8F0", left="E2E8F0", right="E2E8F0")
                            p = cell.paragraphs[0]
                            p.paragraph_format.space_before = Pt(2)
                            p.paragraph_format.space_after = Pt(2)
                            format_inline_text(p, c_text, default_size=Pt(9.0), default_color=BODY_COLOR)

                # Add space after table
                sp = doc.add_paragraph()
                sp.paragraph_format.space_before = Pt(0)
                sp.paragraph_format.space_after = Pt(4)

            in_table = False
            table_rows = []
            # do not continue so current line is processed below

        # 3. Horizontal Rule
        if stripped in ('---', '***', '___'):
            # Add decorative divider paragraph
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(8)
            p_border = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="1" w:color="CBD5E1"/></w:pBdr>')
            p._p.get_or_add_pPr().append(p_border)
            i += 1
            continue

        # 4. Blank Line
        if not stripped:
            i += 1
            continue

        # 5. Headings
        if stripped.startswith('# '):
            text = stripped[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(text)
            run.font.name = "Calibri"
            run.font.size = Pt(22)
            run.bold = True
            run.font.color.rgb = NAVY
            i += 1
            continue

        if stripped.startswith('## '):
            text = stripped[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(text)
            run.font.name = "Calibri"
            run.font.size = Pt(15)
            run.bold = True
            run.font.color.rgb = ACCENT_BLUE
            i += 1
            continue

        if stripped.startswith('### '):
            text = stripped[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            format_inline_text(p, text, default_size=Pt(12.5), default_color=SLATE)
            for r in p.runs:
                r.bold = True
            i += 1
            continue

        if stripped.startswith('#### '):
            text = stripped[5:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            format_inline_text(p, text, default_size=Pt(11.0), default_color=SLATE)
            for r in p.runs:
                r.bold = True
            i += 1
            continue

        # 6. Blockquote
        if stripped.startswith('> '):
            text = stripped[2:].strip()
            table = doc.add_table(rows=1, cols=1)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = table.cell(0, 0)
            set_cell_background(cell, "EFF6FF")
            set_cell_margins(cell, top=100, bottom=100, left=160, right=160)
            set_cell_borders(cell, top="BFDBFE", bottom="BFDBFE", left="2563EB", right="BFDBFE")
            # thick left border for callout
            left_border = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:left w:val="single" w:sz="24" w:space="0" w:color="1E56A0"/><w:top w:val="none"/><w:bottom w:val="none"/><w:right w:val="none"/></w:tcBorders>')
            cell._tc.get_or_add_tcPr().append(left_border)

            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            format_inline_text(p, text, default_size=Pt(10.0), default_color=BODY_COLOR)
            i += 1
            continue

        # 7. Lists (Bullet & Numbered)
        bullet_match = re.match(r'^([\s]*)([\*\-\+])\s+(.*)$', line)
        numbered_match = re.match(r'^([\s]*)(\d+)\.\s+(.*)$', line)

        if bullet_match:
            indent_spaces = len(bullet_match.group(1))
            bullet_text = bullet_match.group(3)
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            if indent_spaces >= 2:
                p.paragraph_format.left_indent = Inches(0.5)
            format_inline_text(p, bullet_text, default_size=Pt(10.5), default_color=BODY_COLOR)
            i += 1
            continue

        if numbered_match:
            num = numbered_match.group(2)
            num_text = numbered_match.group(3)
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.left_indent = Inches(0.25)
            run_num = p.add_run(f"{num}.  ")
            run_num.bold = True
            run_num.font.name = "Calibri"
            run_num.font.size = Pt(10.5)
            run_num.font.color.rgb = ACCENT_BLUE
            format_inline_text(p, num_text, default_size=Pt(10.5), default_color=BODY_COLOR)
            i += 1
            continue

        # 8. Regular Paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        format_inline_text(p, stripped, default_size=Pt(10.5), default_color=BODY_COLOR)
        i += 1

    doc.save(docx_path)
    print(f"Successfully converted {md_path} -> {docx_path}")

if __name__ == "__main__":
    src = Path("docs/DETECTION_AND_PARAMETERS_GUIDE.md")
    dst = Path("docs/DETECTION_AND_PARAMETERS_GUIDE.docx")
    convert_markdown_to_docx(src, dst)
